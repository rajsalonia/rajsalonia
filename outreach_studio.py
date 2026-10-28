"""Deterministic, review-first spreadsheet intake and personalized internship drafts.

This is an explainable assistant, not a trained ML model. It never sends mail,
looks up unknown addresses, or asserts that a company is hiring.
"""
import csv
import io
import re
from email.utils import parseaddr

EMAIL_RE = re.compile(r"(?i)(?<![\w.+-])[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}(?![\w.-])")
EMAIL_OK = re.compile(r"(?i)^[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}$")
ALIASES = {
    "company": {"company", "companyname", "organization", "organisation", "employer", "firm", "startup", "business"},
    "name": {"name", "fullname", "contactname", "hrname", "recruiter", "recruitername", "person", "contactperson"},
    "email": {"email", "emailid", "emailaddress", "e-mail", "hremail", "recruiteremail", "workemail", "contactemail", "mail", "mailid"},
    "role": {"role", "title", "designation", "position", "jobtitle", "contactrole", "department"},
    "notes": {"notes", "remarks", "context", "source", "linkedin", "profile", "website"},
}
BLOCKED_LOCAL = re.compile(r"(?i)^(no.?reply|do.?not.?reply|donotreply|newsletter|billing|abuse|postmaster|mailer.daemon)(?:[+._-]|$)")
PUBLIC_DOMAINS = {"gmail.com", "yahoo.com", "outlook.com", "hotmail.com", "rediffmail.com", "proton.me"}

def norm(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())

def cell(value, limit=500):
    if value is None:
        return ""
    return re.sub(r"[\x00-\x1f\x7f]+", " ", str(value)).strip()[:limit]

def extract_email(value):
    matches = EMAIL_RE.findall(cell(value, 1000))
    return matches[0].lower() if matches else ""

def header_index(values):
    result = {}
    for i, value in enumerate(values):
        normalized = norm(value)
        for key, aliases in ALIASES.items():
            if normalized in {norm(alias) for alias in aliases}:
                result.setdefault(key, i)
    return result

def read_contact_sheet(filename, content, max_rows=6000):
    """Return potential contacts and explicit row warnings, with no writes."""
    if len(content) > 4_000_000:
        raise ValueError("Contact file must be 4 MB or smaller.")
    name = str(filename or "").lower()
    if name.endswith(".csv"):
        try:
            raw = list(csv.reader(io.StringIO(content.decode("utf-8-sig", errors="strict"))))
        except (UnicodeError, csv.Error) as exc:
            raise ValueError("Use a valid UTF-8 CSV file.") from exc
        sheets = [("CSV", raw)]
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheets = []
            for ws in workbook.worksheets[:10]:
                rows = []
                for cells in ws.iter_rows(values_only=True):
                    rows.append(list(cells))
                    if len(rows) >= max_rows + 16:
                        break
                sheets.append((ws.title[:60], rows))
            workbook.close()
        except Exception as exc:
            raise ValueError("Use a valid .xlsx workbook.") from exc
    else:
        raise ValueError("Only CSV and XLSX contact files are supported.")

    found, seen, skipped = [], set(), 0
    for sheet_name, rows in sheets:
        header_row, header = None, {}
        for idx, row in enumerate(rows[:12]):
            detected = header_index(row)
            if "email" in detected or ("company" in detected and "name" in detected):
                header_row, header = idx, detected
                break
        if header_row is None:
            header_row, header = -1, {"company": 0, "name": 1, "email": 2, "role": 3, "notes": 4}
        for row_idx, values in enumerate(rows[header_row + 1:], header_row + 2):
            if len(found) >= max_rows:
                return {"rows": found, "skipped": skipped, "capped": True, "method": "rules-based extraction"}
            if not any(cell(x) for x in values):
                continue
            take = lambda field: cell(values[header[field]], 500) if field in header and len(values) > header[field] else ""
            company = take("company")[:150]
            name = take("name")[:150]
            role = take("role")[:150]
            notes = take("notes")[:500]
            address = extract_email(take("email"))
            if not address:
                for value in values:
                    address = extract_email(value)
                    if address:
                        break
            if address in seen and address:
                skipped += 1
                continue
            if address:
                seen.add(address)
            issues = []
            if not EMAIL_OK.fullmatch(address):
                issues.append("No valid email address found")
            if not company:
                issues.append("Company name needs confirmation")
            local, _, domain = address.partition("@")
            if BLOCKED_LOCAL.search(local):
                issues.append("Automated/no-reply mailbox — excluded from outreach")
            if domain in PUBLIC_DOMAINS:
                issues.append("Personal email domain — verify work purpose and recipient")
            if header_row == -1:
                issues.append("No recognized header — check column mapping")
            ready = bool(address and company and not any("excluded" in issue or "No valid" in issue for issue in issues)
                         and header_row != -1)
            found.append({
                "company": company, "name": name, "email": address, "role": role,
                "notes": notes, "sheet": sheet_name, "row": row_idx,
                "ready": ready, "issues": issues,
            })
    return {"rows": found, "skipped": skipped, "capped": False, "method": "rules-based extraction"}

def classify_contact(contact):
    address = cell(contact.get("email"), 320).lower()
    local, _, domain = address.partition("@")
    role = (cell(contact.get("role")) + " " + local).lower()
    if BLOCKED_LOCAL.search(local):
        category = "do_not_contact"
    elif re.search(r"\b(hr|recruit|talent|hiring|people|career|jobs)\b", role):
        category = "recruiting"
    elif re.search(r"\b(founder|cto|engineering|technical|lead|director)\b", role):
        category = "engineering"
    else:
        category = "unverified"
    return {
        "category": category,
        "recipient_verified": False,
        "advice": ("Do not email an automated mailbox." if category == "do_not_contact" else
                   "Check this person's relationship to hiring and any prior contact before sending."),
        "personal_domain": domain in PUBLIC_DOMAINS,
    }

def compose_draft(contact, *, season="winter", focus="backend"):
    """Use only supported applicant facts, never invented company openings or availability."""
    category = classify_contact(contact)
    if category["category"] == "do_not_contact":
        raise ValueError("Automated/no-reply address cannot be contacted.")
    company = cell(contact.get("company"), 150)
    email = cell(contact.get("email"), 320).lower()
    if not company or not EMAIL_OK.fullmatch(email):
        raise ValueError("Confirm company and email before generating a draft.")
    name = cell(contact.get("name"), 120).split(" ")[0]
    name = re.sub(r"[^A-Za-z'-]", "", name)[:45]
    greeting = ("Hello " + name + ",") if name else "Hello Hiring Team,"
    season = season if season in {"winter", "summer", "general"} else "winter"
    focus = focus if focus in {"backend", "ai", "cloud", "general"} else "backend"
    period = {"winter": "winter 2026–27", "summer": "summer 2027", "general": "upcoming"}[season]
    title = {
        "backend": "Python/backend engineering",
        "ai": "AI engineering",
        "cloud": "cloud/DevOps engineering",
        "general": "software engineering",
    }[focus]
    evidence = {
        "backend": "I work with Python, FastAPI, REST APIs and databases, and previously completed a software engineering internship.",
        "ai": "I have built AI-focused projects and worked with Python, FastAPI and practical software engineering workflows.",
        "cloud": "My interests include AWS, Docker and backend infrastructure, supported by hands-on Python and API projects.",
        "general": "I have software engineering internship experience and have built backend and AI-focused projects.",
    }[focus]
    subject = f"{period.title()} internship enquiry — {title} | Raj Salonia"
    body = (
        f"{greeting}\n\n"
        f"I'm Raj Salonia, a B.Tech Computer Science student graduating in 2028. "
        f"I'm reaching out to ask whether {company} considers students for {period} internships in {title}.\n\n"
        f"{evidence} I would welcome the chance to contribute and learn on a relevant team.\n\n"
        "If there is an appropriate opening or application channel, could you please point me toward it? "
        "I can share my resume for review.\n\n"
        "GitHub: https://github.com/rajsalonia\n"
        "LinkedIn: https://www.linkedin.com/in/rajsalonia/\n\n"
        "Thank you for your time,\nRaj Salonia"
    )
    return {"to": email, "company": company, "subject": subject, "body": body,
            "attach_original_resume": False, "season": season, "focus": focus,
            "model": "rules-based, review required",
            "warning": "Draft only. Verify the recipient, company, eligibility and prior outreach; no email was sent.",
            "classification": category}
